"""Create the editable one-page handout from the skill's canonical guide."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'.claude/skills/agent-issue-authoring/references/guide.md'
OUT=ROOT/'docs/writing-an-agent-issue.docx'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.60)
sec.left_margin=sec.right_margin=Inches(.70)
sec.header_distance=sec.footer_distance=Inches(.22)
for name,size in [('Normal',10.5),('Title',18),('Heading 1',11.5)]:
    st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(4)
    st.paragraph_format.line_spacing=1.03
doc.styles['Title'].paragraph_format.space_after=Pt(5)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(6)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(2)
doc.styles['Heading 1'].font.bold=True
for st in [doc.styles['Title'],doc.styles['Heading 1']]:
    borders=st.element.find('.//'+qn('w:pBdr'))
    if borders is not None:borders.getparent().remove(borders)

def link(p,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'174E73');pr.append(color);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def inline(p,text):
    pattern=r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for bit in re.split(pattern,text):
        if not bit:continue
        if bit.startswith('['):
            m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',bit);link(p,m[1],m[2])
        elif bit.startswith('**'):p.add_run(bit[2:-2]).bold=True
        elif bit.startswith('`'):
            r=p.add_run(bit[1:-1]);r.font.name='Consolas';r.font.size=Pt(9.5)
        else:p.add_run(bit)

for block in SOURCE.read_text().strip().split('\n\n'):
    if block.startswith('# '):
        p=doc.add_paragraph(block[2:],style='Title')
        meta=doc.add_paragraph('AI 1010  October 29 2026');meta.paragraph_format.space_after=Pt(6)
        for r in meta.runs:r.font.size=Pt(9)
    elif block.startswith('## '):doc.add_paragraph(block[3:],style='Heading 1')
    else:
        p=doc.add_paragraph();inline(p,block.replace('\n',' '))
        if block.startswith('**Sources:**'):
            p.paragraph_format.space_before=Pt(3)
            for r in p.runs:r.font.size=Pt(8.5)
f=sec.footer.paragraphs[0]
link(f,'Guide and issue authoring skill','https://github.com/vanderbilt-ai-studies/agent-issue-authoring')
for r in f.runs:r.font.size=Pt(9)
doc.core_properties.title='Write an issue that steers an AI agent'
doc.core_properties.author='Jesse Spencer-Smith'
OUT.parent.mkdir(exist_ok=True);doc.save(OUT)
print(OUT)
